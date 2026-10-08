"""Lecture/écriture du gabarit RUP-seul (`templates/gabarit_rup.xlsx`) --
chantier séparé du 2026-10-07 (voir le plan) : une SEULE feuille "RUP",
structure confirmée en direct sur `Exemplaire.xlsx` (commune Aartselaar,
exemple réel rempli manuellement par l'équipe) :

- Ligne 1 : en-têtes de groupe (fixes) + instructions longues.
- Ligne 2 : sous-en-têtes fixes ("RUP", "Lien d'execution ...") ET,
  à partir de la colonne O, le nom COURT de chaque zone dynamique
  (`svnaam`) -- joue le rôle de `DYNAMIC_RUP_HEADER_ROW` de l'ancien
  gabarit (qui était en ligne 3 là-bas).
- Ligne 3 : "/" pour les colonnes fixes ; pour les colonnes dynamiques,
  le texte légal COMPLET de la zone ("Bestemming...", extrait du PDF
  des prescriptions urbanistiques via `rup_pdf_service.py`) -- remplace
  le simple "/" à fond gris de l'ancien gabarit.
- Ligne 4+ : données, UNE ligne par parcelle RETENUE (filtrée : au
  moins un RUP région/province/commune applicable -- voir main_rup.py).

Colonne A ("Nom de RUP", SANS en-tête texte -- confirmé en direct sur
l'Exemplaire, la colonne n'a aucun libellé) : remplie UNE SEULE FOIS par
groupe de lignes consécutives partageant le même RUP (`fichelink`),
cellule fusionnée sur tout le groupe -- reproduit exactement la
convention observée (ex. "Omgeving Domein Solhof" fusionné sur A4:A15).

Toutes les autres colonnes fixes (B->N) et la mécanique de colonnes
dynamiques (O+) reprennent le même principe déjà éprouvé dans
`excel_service.py` (gabarit à 190 colonnes), avec seulement les indices
de ligne/colonne adaptés à cette structure à 3 lignes d'en-tête au lieu
de 4."""

from __future__ import annotations

from pathlib import Path
from typing import Dict, Optional, Set, Tuple

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, PatternFill
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from services.excel_service import couleur_depuis_legende, vers_capakey_court  # noqa: F401 -- réexporté pour main_rup.py

FIRST_DATA_ROW = 4
NOM_FEUILLE = "RUP"

COL_NOM_RUP = 1  # A -- "Nom de RUP", fusionné par groupe
COL_COMMUNE = 2  # B
COL_CODE_POSTAL = 3  # C
COL_RUE = 4  # D
COL_NUMERO = 5  # E
COL_NUMERO_CADASTRAL = 6  # F

# Colonnes (lien, texte) par niveau -- voir le docstring du module et
# wfs_rup_service.py. Colonnes 9 (I) et 12 (L) sont de simples
# séparateurs visuels, toujours "/" (voir Exemplaire.xlsx), jamais de
# donnée réelle.
_RUP_COLONNES = {"region": (7, 8), "province": (10, 11), "commune": (13, 14)}
_COLONNES_SEPARATRICES = (9, 12)

DYNAMIC_RUP_FIRST_COL = 15  # O
DYNAMIC_RUP_HEADER_ROW = 2  # svnaam (nom court de la zone)
DYNAMIC_RUP_BESTEMMING_ROW = 3  # texte légal complet ("Bestemming...")


def charger_classeur(excel_path: Path) -> Workbook:
    return openpyxl.load_workbook(excel_path)


def feuille_rup(wb: Workbook) -> Worksheet:
    return wb[NOM_FEUILLE]


def trouver_premiere_ligne_vide(ws: Worksheet) -> int:
    """Première ligne (>= FIRST_DATA_ROW) où la colonne F (Numéro
    cadastral) est vide -- jamais la colonne A, qui est délibérément
    vide sur la plupart des lignes (fusion par groupe, voir le
    docstring du module)."""
    r = FIRST_DATA_ROW
    while ws.cell(row=r, column=COL_NUMERO_CADASTRAL).value not in (None, ""):
        r += 1
    return r


def lire_capakeys_deja_ecrits(ws: Worksheet) -> Set[str]:
    capakeys: Set[str] = set()
    derniere = trouver_premiere_ligne_vide(ws) - 1
    for r in range(FIRST_DATA_ROW, derniere + 1):
        v = ws.cell(row=r, column=COL_NUMERO_CADASTRAL).value
        if v not in (None, ""):
            capakeys.add(str(v).strip())
    return capakeys


def lire_cle_groupe(ws: Worksheet, row: int) -> Optional[Tuple[str, str, str]]:
    """Identité RUP de `row` (texte région/province/commune, ex.
    `("/", "/", "RUP_11001_214_00005_00001")`) -- DEUX lignes
    consécutives appartiennent au même groupe visuel ("Nom de RUP") ssi
    cette identité est RIGOUREUSEMENT identique (voir main_rup.py).
    `None` si `row` est avant la première ligne de données (fichier
    encore vide)."""
    if row < FIRST_DATA_ROW:
        return None
    return tuple(ws.cell(row=row, column=_RUP_COLONNES[niveau][1]).value for niveau in ("region", "province", "commune"))


def debut_groupe_existant(ws: Worksheet, row: int) -> int:
    """Première ligne du groupe (fusion colonne A) auquel `row`
    appartient déjà -- `row` lui-même si cette ligne n'est pas (encore)
    fusionnée (groupe d'une seule ligne pour l'instant)."""
    for plage in ws.merged_cells.ranges:
        if plage.min_col == COL_NOM_RUP and plage.max_col == COL_NOM_RUP and plage.min_row <= row <= plage.max_row:
            return plage.min_row
    return row


def etendre_groupe_nom_rup(ws: Worksheet, row_debut: int, row_fin: int) -> None:
    """Étend la fusion de la colonne A (groupe RUP courant) jusqu'à
    `row_fin` -- défusionne d'abord la plage existante le cas échéant
    (sinon openpyxl refuse une fusion qui chevauche une fusion déjà en
    place), nécessaire pour continuer un groupe à travers une reprise
    de run (voir main_rup.py). Ne fait rien si `row_fin == row_debut`
    (rien à étendre encore)."""
    if row_fin <= row_debut:
        return
    for plage in list(ws.merged_cells.ranges):
        if plage.min_col == COL_NOM_RUP and plage.max_col == COL_NOM_RUP and plage.min_row == row_debut:
            ws.unmerge_cells(str(plage))
            break
    ws.merge_cells(start_row=row_debut, start_column=COL_NOM_RUP, end_row=row_fin, end_column=COL_NOM_RUP)


def ecrire_identite(
    ws: Worksheet, row: int, *, commune: str, code_postal: str, rue: str, numero: str, capakey: str,
) -> None:
    ws.cell(row=row, column=COL_COMMUNE, value=commune)
    ws.cell(row=row, column=COL_CODE_POSTAL, value=code_postal)
    ws.cell(row=row, column=COL_RUE, value=rue)
    ws.cell(row=row, column=COL_NUMERO, value=numero)
    ws.cell(row=row, column=COL_NUMERO_CADASTRAL, value=capakey)
    for col in _COLONNES_SEPARATRICES:
        ws.cell(row=row, column=col, value="/")


def ecrire_rup(ws: Worksheet, row: int, niveau: str, *, lien: str, texte: str) -> None:
    col_lien, col_texte = _RUP_COLONNES[niveau]
    ws.cell(row=row, column=col_lien, value=lien)
    ws.cell(row=row, column=col_texte, value=texte)


def ecrire_nom_rup(ws: Worksheet, row: int, nom: str) -> None:
    """Écrit `nom` (ex. "Omgeving Domein Solhof") en colonne A de
    `row` -- PREMIÈRE ligne d'un nouveau groupe RUP (voir
    `lire_cle_groupe`/`etendre_groupe_nom_rup`, appelés par
    main_rup.py pour décider si une ligne démarre un nouveau groupe ou
    continue le précédent). Pas de fusion ici : une ligne qui démarre
    seule n'en a pas encore besoin, `etendre_groupe_nom_rup` s'en charge
    dès qu'une deuxième ligne rejoint le groupe."""
    cell = ws.cell(row=row, column=COL_NOM_RUP, value=nom)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---------------------------------------------------------------------------
# Colonnes dynamiques (une par `svnaam` rencontré) -- même principe que
# `excel_service.py::trouver_ou_creer_colonne_dynamique_rup`, adapté à
# cette structure (en-tête ligne 2, texte Bestemming ligne 3 au lieu du
# simple "/" à fond gris).
# ---------------------------------------------------------------------------

def index_colonnes_dynamiques(ws: Worksheet) -> Dict[str, int]:
    index: Dict[str, int] = {}
    for col in range(DYNAMIC_RUP_FIRST_COL, ws.max_column + 1):
        v = ws.cell(row=DYNAMIC_RUP_HEADER_ROW, column=col).value
        if v:
            index[str(v)] = col
    return index


def trouver_ou_creer_colonne_dynamique(
    ws: Worksheet, index: Dict[str, int], svnaam: str, bestemming: str,
    *, couleur_fond: Optional[str] = None, chemin_image_motif: Optional[Path] = None,
) -> int:
    """Renvoie l'index de colonne pour `svnaam`, la créant si besoin --
    en-tête (ligne 2) + texte légal complet `bestemming` (ligne 3,
    extrait du PDF via `rup_pdf_service.py`, voir main_rup.py).

    `couleur_fond` (hex SANS '#') et `chemin_image_motif` (chemin d'une
    petite image PNG) sont mutuellement exclusifs -- voir
    `services/rup_legende_service.py` et son appelant (main_rup.py) :
    couleur unie -> fond de cellule direct ; motif (hachures...) ->
    petite image fidèle au vrai motif officiel, ancrée sur la cellule.
    Aucun des deux si la couleur réelle n'a pas pu être déterminée
    (jamais de couleur devinée). Met `index` à jour en place. N'écrase
    JAMAIS `bestemming` si la colonne existe déjà (toujours le même
    texte pour un même `svnaam`, premier arrivé suffit -- évite de
    retélécharger/réécrire inutilement)."""
    if svnaam in index:
        return index[svnaam]
    nouvelle_col = max([DYNAMIC_RUP_FIRST_COL - 1, *index.values()]) + 1
    cell = ws.cell(row=DYNAMIC_RUP_HEADER_ROW, column=nouvelle_col, value=svnaam)
    cell.font = cell.font.copy(bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    if couleur_fond:
        cell.fill = PatternFill(start_color=couleur_fond, end_color=couleur_fond, fill_type="solid")
    elif chemin_image_motif is not None:
        image = XLImage(str(chemin_image_motif))
        image.width = image.height = 18
        ws.add_image(image, cell.coordinate)
    texte_cell = ws.cell(row=DYNAMIC_RUP_BESTEMMING_ROW, column=nouvelle_col, value=(bestemming or "/"))
    texte_cell.alignment = Alignment(wrap_text=True, vertical="top")
    index[svnaam] = nouvelle_col
    return nouvelle_col


def ecrire_zones_dynamiques(ws: Worksheet, row: int, index: Dict[str, int], colonnes_gagnantes: Set[int]) -> None:
    for col in index.values():
        ws.cell(row=row, column=col, value=("O" if col in colonnes_gagnantes else "N"))
