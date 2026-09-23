"""Réordonnance un fichier d'état DÉJÀ traité pour que chaque rue soit un
bloc contigu, et que ses lignes soient dans l'ordre le long de la route
(un côté puis l'autre, dans l'ordre d'apparition -- même convention que
`decouverte_service.py::_ordonner_le_long`), avec ou sans numéro de maison.

**Pourquoi ce module existe séparément de `main.py`** : `executer_traitement`
écrit toujours une nouvelle parcelle à la première ligne vide du fichier
(voir `excel_service.py::trouver_premiere_ligne_vide`) -- correct pour UN
run isolé (les parcelles d'une même rue, dans le même run, sont déjà
ordonnées AVANT d'être écrites, voir `_ordonner_le_long`), mais PAS entre
deux runs séparés sur la MÊME rue : les parcelles découvertes au 2e run
(ex. les sœurs sans adresse ou les parcelles trouvées par géométrie, ajoutées
au pipeline après un premier run classique) s'ajoutent à la SUITE de tout
le fichier, après toutes les autres rues -- écart réel confirmé le
2026-09-22 sur Sint-Gillis-Waas. Réécrire `executer_traitement` pour insérer
directement à la bonne position casserait la garantie de robustesse "chaque
parcelle sauvegardée immédiatement, reprise sûre après un crash" (un insert
en cours de route décale TOUTES les lignes suivantes, y compris celles
d'une rue pas encore traitée dans CE run) -- ce module fait donc exprès une
passe SÉPARÉE, après coup, sur un fichier déjà stable.

Principe : pour chaque rue déjà présente dans le fichier (dans son ordre de
première apparition, jamais changé), on la redécouvre avec le pipeline
COMPLET actuel (adresses + parcelles sœurs + géométrie -- donc toujours
LA MÊME LISTE de parcelles qu'un `executer_traitement` referait aujourd'hui
sur cette rue), on en déduit l'ordre "autorité", puis on trie les lignes
DÉJÀ PRÉSENTES du fichier selon cet ordre. Ne SUPPRIME ni n'AJOUTE jamais
de ligne -- purement une permutation. Une ligne dont la parcelle n'est plus
retrouvée par la redécouverte (rare -- parcelle fusionnée/renumérotée côté
cadastre) est gardée, en dernier dans le bloc de sa rue, JAMAIS perdue."""

from __future__ import annotations

import shutil
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from services.adressen_service import AdressenService
from services.cadastre_service import CadastreService
from services.decouverte_geometrique_service import DecouverteGeometrique
from services.decouverte_service import decouvrir_parcelles
from services.excel_service import (
    COL_NUMERO_CADASTRAL, COL_RUE, FIRST_DATA_ROW, charger_classeur, feuille_principale, feuille_rup, sauvegarder,
    vers_capakey_court,
)
from utils.logger import get_logger

_logger = get_logger("services.reordonnancement_service")


def _derniere_ligne(ws: Worksheet) -> int:
    r = FIRST_DATA_ROW
    while ws.cell(row=r, column=1).value not in (None, ""):
        r += 1
    return r - 1


def _rues_par_ordre_apparition(ws: Worksheet, derniere: int) -> List[str]:
    vues: List[str] = []
    connues: set = set()
    for r in range(FIRST_DATA_ROW, derniere + 1):
        rue = ws.cell(row=r, column=COL_RUE).value
        if rue not in connues:
            connues.add(rue)
            vues.append(rue)
    return vues


def _copier_ligne(ws: Worksheet, r: int) -> List[Tuple[object, object]]:
    return [(c.value, copy(c._style)) for c in ws[r][: ws.max_column]]


def _ecrire_ligne(ws: Worksheet, r: int, contenu: List[Tuple[object, object]]) -> None:
    for col, (valeur, style) in enumerate(contenu, start=1):
        cell = ws.cell(row=r, column=col)
        cell.value = valeur
        cell._style = copy(style)


def reordonner_fichier(
    excel_path: Path, commune: str, adressen: AdressenService, cadastre: CadastreService,
    geometrique: DecouverteGeometrique,
) -> Tuple[int, int, int]:
    """Réordonne `excel_path` en place (feuille principale ET RUP, même
    permutation -- les deux restent alignées ligne à ligne). Sauvegarde
    horodatée AVANT toute écriture (même discipline que partout ailleurs
    dans ce projet pour un patch direct sur un fichier déjà traité).

    Renvoie `(n_lignes_deplacees, n_rues, n_lignes_non_retrouvees)`."""
    backup = excel_path.with_name(
        f"{excel_path.stem}.backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}-avant-reordonnancement.xlsx",
    )
    shutil.copyfile(excel_path, backup)
    _logger.info("Sauvegarde avant réordonnancement : %s", backup.name)

    wb = charger_classeur(excel_path)
    principale = feuille_principale(wb)
    rup = feuille_rup(wb)
    derniere = _derniere_ligne(principale)
    lignes = list(range(FIRST_DATA_ROW, derniere + 1))

    if rup is not None:
        for r in lignes:
            a = (principale.cell(row=r, column=COL_RUE).value, principale.cell(row=r, column=COL_NUMERO_CADASTRAL).value)
            b = (rup.cell(row=r, column=COL_RUE).value, rup.cell(row=r, column=COL_NUMERO_CADASTRAL).value)
            if a != b:
                raise ValueError(
                    f"Feuille principale et RUP désynchronisées à la ligne {r} ({a} != {b}) -- "
                    "réordonnancement annulé AVANT toute écriture, fichier intact.",
                )

    rues = _rues_par_ordre_apparition(principale, derniere)
    _logger.info("Réordonnancement de '%s' : %d rue(s), %d ligne(s).", excel_path.name, len(rues), len(lignes))

    nouvel_ordre: List[int] = []
    n_non_retrouvees = 0
    for rue in rues:
        lignes_rue = [r for r in lignes if principale.cell(row=r, column=COL_RUE).value == rue]
        try:
            parcelles = decouvrir_parcelles(commune, rue, adressen, cadastre, geometrique)
        except Exception as exc:  # noqa: BLE001 -- une rue en échec ne doit jamais faire perdre l'ordre des autres
            _logger.warning(
                "Redécouverte de '%s' a échoué (%s: %s) -- ordre de cette rue INCHANGÉ (déjà tel quel dans le "
                "fichier), les autres rues sont quand même réordonnées.", rue, type(exc).__name__, exc,
            )
            nouvel_ordre.extend(lignes_rue)
            continue
        rang: Dict[str, int] = {vers_capakey_court(p.parcelle.reference): i for i, p in enumerate(parcelles)}
        n_inconnues_ici = sum(
            1 for r in lignes_rue if str(principale.cell(row=r, column=COL_NUMERO_CADASTRAL).value).strip() not in rang
        )
        n_non_retrouvees += n_inconnues_ici
        if n_inconnues_ici:
            _logger.warning(
                "'%s' : %d ligne(s) du fichier non retrouvée(s) par la redécouverte -- gardée(s) en fin de "
                "bloc, jamais supprimée(s).", rue, n_inconnues_ici,
            )
        lignes_rue_triees = sorted(
            lignes_rue,
            key=lambda r: rang.get(str(principale.cell(row=r, column=COL_NUMERO_CADASTRAL).value).strip(), len(rang)),
        )
        nouvel_ordre.extend(lignes_rue_triees)

    assert sorted(nouvel_ordre) == lignes, "permutation invalide (ligne perdue ou dupliquée) -- jamais écrit"

    feuilles = [principale] + ([rup] if rup is not None else [])
    for ws in feuilles:
        source = {r: _copier_ligne(ws, r) for r in lignes}
        for cible, src in zip(lignes, nouvel_ordre):
            _ecrire_ligne(ws, cible, source[src])

    n_deplacees = sum(1 for a, b in zip(lignes, nouvel_ordre) if a != b)
    sauvegarder(wb, excel_path)
    _logger.info(
        "Réordonnancement de '%s' terminé : %d/%d ligne(s) déplacée(s), %d ligne(s) non retrouvée(s) "
        "(gardées, jamais perdues).", excel_path.name, n_deplacees, len(lignes), n_non_retrouvees,
    )
    return n_deplacees, len(rues), n_non_retrouvees
